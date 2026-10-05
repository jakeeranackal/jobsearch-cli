"""Read resumes from txt/md/docx/pdf, model their structure, write docx/markdown.

The structure is deliberately loose: a resume is a list of sections, each a
list of items that are headings (a job/school line), bullets, or plain text.
That's enough to reorder and rewrite bullets without understanding layouts.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SECTION_NAMES = (
    "summary", "professional summary", "profile", "objective", "about",
    "experience", "work experience", "professional experience", "employment",
    "employment history", "relevant experience", "work history", "career history",
    "education", "skills", "technical skills", "core skills", "key skills", "core competencies",
    "skills & tools", "skills and tools", "tools", "projects", "selected projects",
    "certifications", "licenses & certifications", "certificates", "awards", "honors",
    "leadership", "leadership experience", "volunteer", "volunteering", "publications",
    "activities", "interests", "languages", "additional information", "additional",
)
_BULLET = re.compile(r"^\s*[\-\*•●▪–➢►‣o]\s+")
_SECTION_WORD = re.compile(
    r"experience|education|skill|project|summary|certif|award|leadership|volunteer|"
    r"language|interest|activit|publication|profile|objective|competenc|employment|"
    r"history|training|course|honor|tools"
)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")


@dataclass
class Item:
    kind: str  # heading | bullet | text
    text: str


@dataclass
class Section:
    title: str
    items: list[Item] = field(default_factory=list)

    @property
    def key(self) -> str:
        t = self.title.lower()
        if "skill" in t or "competenc" in t or t == "tools":
            return "skills"
        if "experience" in t or "employment" in t or "history" in t:
            return "experience"
        if t in ("summary", "professional summary", "profile", "objective", "about"):
            return "summary"
        return t


@dataclass
class Resume:
    header: list[str] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)

    def section(self, key: str) -> Section | None:
        return next((s for s in self.sections if s.key == key), None)

    def bullets(self) -> list[str]:
        return [i.text for s in self.sections for i in s.items if i.kind == "bullet"]

    def to_text(self) -> str:
        out = list(self.header)
        for s in self.sections:
            out += ["", s.title.upper()]
            for i in s.items:
                out.append(f"- {i.text}" if i.kind == "bullet" else i.text)
        return "\n".join(out).strip() + "\n"


def read_text(path: str | Path) -> str:
    """Plain text from a resume file. docx list paragraphs become '- ' bullets."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Resume not found: {p}")
    suffix = p.suffix.lower()
    if suffix == ".docx":
        import docx

        d = docx.Document(str(p))
        lines = []
        for para in d.paragraphs:
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            is_list = "list" in style or para._p.pPr is not None and para._p.pPr.numPr is not None
            lines.append(f"- {text}" if is_list and not _BULLET.match(text) else text)
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
        return "\n".join(lines)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(p))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    text = p.read_text(encoding="utf-8", errors="replace")
    return md_to_text(text) if suffix in (".md", ".markdown") else text


def md_to_text(md: str) -> str:
    """Markdown resume to the plain layout parse() expects.

    '# Name' stays the header, '## Section' becomes a section title, '### Role'
    and '**Role**' become role headings, and bullets pass through.
    """
    lines = []
    for ln in md.splitlines():
        s = ln.strip()
        if s.startswith("## ") and not s.startswith("### "):
            lines += ["", s[3:].strip().upper()]
        elif s.startswith("#"):
            lines.append(s.lstrip("#").strip())
        elif s.startswith("**") and s.endswith("**") and len(s) > 4:
            lines.append(s.strip("*").strip())
        elif s.startswith(("- ", "* ", "+ ")):
            lines.append(ln.replace("**", ""))
        else:
            lines.append(ln.replace("**", "").replace("__", ""))
    return "\n".join(lines)


def _is_section_title(line: str) -> bool:
    s = line.strip().rstrip(":").strip()
    if not s or len(s) > 40:
        return False
    if s.lower() in SECTION_NAMES:
        return True
    # All-caps lines count only if they look like a section ("RELEVANT EXPERIENCE"),
    # so a name or company in caps isn't mistaken for one.
    return s.upper() == s and len(s.split()) <= 4 and bool(_SECTION_WORD.search(s.lower()))


def parse(text: str) -> Resume:
    resume = Resume()
    current: Section | None = None
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if _is_section_title(line):
            current = Section(title=line.strip().rstrip(":").title())
            resume.sections.append(current)
            continue
        if current is None:
            resume.header.append(line.strip())
            continue
        if _BULLET.match(line):
            current.items.append(Item("bullet", _BULLET.sub("", line).strip()))
        elif current.items and current.items[-1].kind == "bullet" and raw[:1] in (" ", "\t"):
            # Wrapped bullet continuation
            current.items[-1].text += " " + line.strip()
        else:
            kind = "text" if current.key in ("summary", "skills") else "heading"
            current.items.append(Item(kind, line.strip()))
    return resume


def load(path: str | Path) -> Resume:
    return parse(read_text(path))


def write_markdown(resume: Resume, path: Path) -> Path:
    lines = []
    if resume.header:
        lines.append(f"# {resume.header[0]}")
        lines += resume.header[1:]
    for s in resume.sections:
        lines += ["", f"## {s.title}"]
        for i in s.items:
            if i.kind == "bullet":
                lines.append(f"- {i.text}")
            elif i.kind == "heading":
                lines += ["", f"**{i.text}**"]
            else:
                lines.append(i.text)
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def write_docx(resume: Resume, path: Path) -> Path:
    """Single-column, no tables or text boxes, standard headings: ATS-safe."""
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Pt(54)
        sec.top_margin = sec.bottom_margin = Pt(46)

    if resume.header:
        name = d.add_paragraph()
        run = name.add_run(resume.header[0])
        run.bold = True
        run.font.size = Pt(16)
        if len(resume.header) > 1:
            d.add_paragraph(" | ".join(resume.header[1:]))
    for s in resume.sections:
        h = d.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        run = h.add_run(s.title.upper())
        run.bold = True
        run.font.size = Pt(11.5)
        for i in s.items:
            if i.kind == "bullet":
                p = d.add_paragraph(i.text, style="List Bullet")
            elif i.kind == "heading":
                p = d.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.add_run(i.text).bold = True
            else:
                p = d.add_paragraph(i.text)
            p.paragraph_format.space_after = Pt(2)
    d.save(str(path))
    return path


def ats_check(path: str | Path) -> list[tuple[str, str]]:
    """Return (level, message) findings about how an ATS will read this file."""
    p = Path(path)
    findings: list[tuple[str, str]] = []
    suffix = p.suffix.lower()
    if suffix not in (".docx", ".pdf", ".txt", ".md"):
        findings.append(("bad", f"{suffix} files are often rejected. Use .docx or .pdf."))

    if suffix == ".docx":
        import docx

        d = docx.Document(str(p))
        if d.tables:
            findings.append(("warn", f"{len(d.tables)} table(s). Many ATSs scramble or drop table text."))
        header_text = " ".join(
            par.text for sec in d.sections for part in (sec.header, sec.footer)
            for par in part.paragraphs
        ).strip()
        if header_text:
            findings.append(("warn", "Text in the page header/footer. Some ATSs skip it, so keep contact info in the body."))
        xml = d.element.xml
        if "txbxContent" in xml:
            findings.append(("bad", "Text boxes found. ATSs usually can't read text inside them."))
        if re.search(r'<w:cols [^>]*w:num="[2-9]"', xml):
            findings.append(("warn", "Multi-column layout. Columns often get read in the wrong order."))
        if d.inline_shapes:
            findings.append(("warn", "Images/icons found. They're invisible to an ATS; don't put info in them."))
    text = read_text(p)
    if suffix == ".pdf" and len(text.strip()) < 200:
        findings.append(("bad", "Almost no text extracted. This PDF is likely a scanned image; ATSs will see a blank page."))

    words = len(text.split())
    if not EMAIL_RE.search(text):
        findings.append(("bad", "No email address found in the text."))
    if not PHONE_RE.search(text):
        findings.append(("warn", "No phone number found in the text."))
    resume = parse(text)
    keys = {s.key for s in resume.sections}
    for needed in ("experience", "education", "skills"):
        if needed not in keys:
            findings.append(("warn", f"No standard '{needed.title()}' heading detected. Use plain headings ATSs recognize."))
    if words > 1100:
        findings.append(("warn", f"{words} words. Over ~2 pages; recruiters skim, trim older roles."))
    elif words < 250:
        findings.append(("warn", f"Only {words} words extracted. Check that nothing was lost."))
    odd = set(re.findall(r"[☀-➿-]", text))
    if odd:
        findings.append(("warn", f"Unusual symbols {' '.join(sorted(odd))} may show up as garbage."))
    if not findings:
        findings.append(("ok", "No ATS problems found."))
    return findings
