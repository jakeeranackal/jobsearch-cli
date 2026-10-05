import docx

from jobsearch import resume_io


def test_parse_sections_and_bullets(resume_text):
    r = resume_io.parse(resume_text)
    assert r.header[0] == "Jordan Lee"
    keys = [s.key for s in r.sections]
    assert keys[:3] == ["summary", "experience", "skills"]
    assert len(r.bullets()) == 4
    exp = r.section("experience")
    assert exp.items[0].kind == "heading"


def test_name_in_caps_is_header_not_section():
    r = resume_io.parse("JORDAN LEE\njordan@x.com\n\nEXPERIENCE\nACME CORP\n- Did things\n")
    assert r.header == ["JORDAN LEE", "jordan@x.com"]
    assert [s.title for s in r.sections] == ["Experience"]
    assert r.sections[0].items[0].text == "ACME CORP"


def test_docx_roundtrip(tmp_path, resume_text):
    r = resume_io.parse(resume_text)
    path = resume_io.write_docx(r, tmp_path / "out.docx")
    text = resume_io.read_text(path)
    assert "- Automated a manual reconciliation" in text
    again = resume_io.parse(text)
    assert len(again.bullets()) == len(r.bullets())


def test_ats_check_flags_tables(tmp_path, resume_text):
    d = docx.Document()
    for line in resume_text.splitlines():
        d.add_paragraph(line)
    d.add_table(rows=1, cols=2).cell(0, 0).text = "SQL"
    path = tmp_path / "r.docx"
    d.save(path)
    messages = " ".join(m for _, m in resume_io.ats_check(path))
    assert "table" in messages


def test_ats_check_missing_email(tmp_path):
    path = tmp_path / "r.txt"
    path.write_text("Name\n\nEXPERIENCE\n- Did a thing\n", encoding="utf-8")
    findings = resume_io.ats_check(path)
    assert ("bad", "No email address found in the text.") in findings
